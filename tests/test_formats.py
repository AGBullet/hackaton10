import uuid
import fitz,pytest
from inspector.config import ROOT
from inspector.extraction import document_pdf
from inspector.registry import create_object,register

@pytest.mark.parametrize('kind',['docx','xml','xlsx'])
def test_office_and_xml_preview(isolated_db,kind):
    dest=ROOT/'data/test-fixtures'/uuid.uuid4().hex;dest.mkdir(parents=True)
    p=dest/('source.'+kind)
    if kind=='docx':
        from docx import Document
        doc=Document();doc.add_paragraph('Площадь застройки: 1234 м²');doc.save(p)
    elif kind=='xlsx':
        from openpyxl import Workbook
        doc=Workbook();doc.active.append(['Площадь застройки','1234','м²']);doc.save(p)
    else:p.write_text('<document><parameter>Площадь застройки: 1234 м²</parameter></document>',encoding='utf-8')
    obj=create_object('FORMAT-'+uuid.uuid4().hex);d=register(p,obj,{'stage':'PD'})
    with fitz.open(document_pdf(d)) as pdf:
        assert '1234' in ''.join(page.get_text() for page in pdf)

def test_upload_and_manifest_validation(isolated_db):
    from fastapi.testclient import TestClient
    from inspector.app import app
    from inspector.db import one
    obj=create_object('BAD-UPLOAD-'+uuid.uuid4().hex)
    with TestClient(app) as c:
        r=c.post('/api/objects/'+obj+'/upload',data={'stage':'PD'},files={'files':('broken.pdf',b'not a pdf','application/pdf')})
        assert r.status_code==400
        assert one('SELECT count(*) AS n FROM documents WHERE object_id=%s',(obj,))['n']==0
        r=c.post('/api/objects/'+obj+'/manifest',json={'documents':[{'stage':'PD','document_code':'PZ','gold_answer':'123'}],'reason':'Test manifest'})
        assert r.status_code==400
