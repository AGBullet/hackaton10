import collections, hashlib, io, json, math, os, re, time, unicodedata, uuid
from functools import lru_cache
from contextlib import nullcontext
from pathlib import Path
import fitz
from PIL import Image
import pytesseract
from docx import Document
import openpyxl
from defusedxml import ElementTree
from .config import ROOT, DATA, PARSER_VERSION, inside_root
from .db import execute,one,query
from .registry import uid

pytesseract.pytesseract.tesseract_cmd=os.getenv('TESSERACT_CMD',r'C:\Program Files\Tesseract-OCR\tesseract.exe')

def norm(text):
    return re.sub(r'\s+',' ',unicodedata.normalize('NFC',str(text))).strip()

def native_text_usable(text):
    controls=sum((ord(c)<32 and c not in '\n\r\t') or 127<=ord(c)<=159 or unicodedata.category(c)=='Co' for c in text)
    return (sum(c.isalnum() for c in text)>=40 and text.count('\ufffd')<max(2,len(text)*.02)
            and controls<max(4,len(text)*.01))

def bbox_norm(rect,width,height):
    r=list(rect)
    return [round(max(0,min(1,r[i]/(width if i%2==0 else height))),6) for i in range(4)]

def render_region(page,box,max_side=1024,scale_limit=4):
    """Render a visible rotated-page region directly, retaining small drawing labels."""
    if len(box)!=4 or not all(math.isfinite(v) and 0<=v<=1 for v in box) or box[0]>=box[2] or box[1]>=box[3]:raise ValueError('Неверная область')
    rect=fitz.Rect(box[0]*page.rect.width,box[1]*page.rect.height,box[2]*page.rect.width,box[3]*page.rect.height)
    factor=min(scale_limit,max_side/max(rect.width,rect.height))
    return page.get_displaylist().get_pixmap(matrix=fitz.Matrix(factor,factor),clip=rect,alpha=False)

def pdf_lines(page):
    lines=[]
    for block in page.get_text('dict',sort=True)['blocks']:
        for line in block.get('lines',[]):
            text=''.join(s['text'] for s in line['spans']).replace('\x00','').strip()
            if not text:continue
            rect=fitz.Rect(line['bbox'])*page.rotation_matrix
            rect=rect & page.rect
            if rect.is_empty:continue
            lines.append({'text':text,'bbox':bbox_norm(rect,page.rect.width,page.rect.height),'confidence':1.0})
    return lines

def write_cache(cache,result):
    cache.parent.mkdir(parents=True,exist_ok=True)
    tmp=cache.with_name(cache.name+'.'+uuid.uuid4().hex+'.part')
    tmp.write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
    try:
        for attempt in range(6):
            try:
                tmp.replace(cache)
                break
            except PermissionError:
                if attempt==5:raise
                time.sleep(.05*(attempt+1))
    finally:
        tmp.unlink(missing_ok=True)


@lru_cache(maxsize=1)
def ocr_engine_version():
    return str(pytesseract.get_tesseract_version())


def ocr_image(img,dpi):
    # Identical rendered pages recur inside different real certificate bundles.
    # Reuse only exact pixels/settings, never perceptual similarity or a title.
    identity=json.dumps([ocr_engine_version(),'rus+eng','psm11',dpi,img.mode,img.size]).encode()
    digest=hashlib.sha256(identity+img.tobytes()).hexdigest()
    cache=DATA/'cache'/'ocr-pixels-v1'/digest[:2]/(digest+'.json')
    try:
        stored=json.loads(cache.read_text(encoding='utf-8'))
        if isinstance(stored,list) and all(isinstance(l,dict) and {'text','bbox','confidence'}<=l.keys() for l in stored):
            img.info['inspector_ocr_cache_hit']=True
            return stored
    except (OSError,ValueError):pass
    img.info['inspector_ocr_cache_hit']=False
    data=pytesseract.image_to_data(img,lang='rus+eng',config=f'--psm 11 --dpi {dpi}',output_type=pytesseract.Output.DICT,timeout=120)
    grouped=collections.defaultdict(list)
    for i,txt in enumerate(data['text']):
        if txt.strip():grouped[(data['block_num'][i],data['par_num'][i],data['line_num'][i])].append(i)
    lines=[]
    for indexes in grouped.values():
        x=min(data['left'][i] for i in indexes); y=min(data['top'][i] for i in indexes)
        x1=max(data['left'][i]+data['width'][i] for i in indexes); y1=max(data['top'][i]+data['height'][i] for i in indexes)
        lines.append({'text':' '.join(data['text'][i] for i in indexes),'bbox':bbox_norm([x,y,x1,y1],img.width,img.height),'confidence':max(0,sum(float(data['conf'][i]) for i in indexes)/len(indexes)/100)})
    write_cache(cache,lines)
    return lines

def extract_page(path,page_number,allow_ocr=True,document=None,collect_geometry=False):
    started=time.perf_counter()
    with (nullcontext(document) if document is not None else fitz.open(path)) as pdf:
        page=pdf[page_number-1]
        lines=pdf_lines(page)
        txt='\n'.join(l['text'] for l in lines)
        good=native_text_usable(txt)
        method='text';dpi=None;ocr_cache_hit=False
        image_count=len(page.get_images())
        if not good and allow_ocr:
            dpi=int(os.getenv('OCR_DPI','200'))
            max_pixels=int(os.getenv('OCR_MAX_PIXELS','18000000'))
            dpi=min(dpi,max(50,int(72*math.sqrt(max_pixels/(page.rect.width*page.rect.height)))))
            pix=page.get_pixmap(dpi=dpi,alpha=False)
            img=Image.open(io.BytesIO(pix.tobytes('png')))
            lines=ocr_image(img,dpi)
            ocr_cache_hit=img.info.get('inspector_ocr_cache_hit',False)
            txt='\n'.join(l['text'] for l in lines);method='tesseract-rus-eng'
        conf=sum(l['confidence'] for l in lines)/max(1,len(lines))
        quality='GOOD' if len(txt.strip())>=40 and conf>=.80 and (method!='text' or good) else ('OCR_REQUIRED' if not allow_ocr else 'LOW_QUALITY')
        drawings=None
        # The dimension tool reads vectors on demand. Eagerly walking every CAD
        # path on every corpus page costs seconds and is not a matrix check.
        if collect_geometry and method=='text' and path.stat().st_size<100*1024**2:
            try:
                drawings=[{'rect':bbox_norm(d['rect']*page.rotation_matrix,page.rect.width,page.rect.height),'items':len(d['items'])} for d in page.get_drawings()[:1500]]
            except Exception:pass
        content={'lines':lines,'width':page.rect.width,'height':page.rect.height,'rotation':page.rotation,'cropbox':list(page.cropbox),'mediabox':list(page.mediabox),'dpi':dpi,'image_count':image_count,'vector_paths':len(drawings) if drawings is not None else None,'geometry':(drawings or [])[:200],'geometry_status':'COLLECTED' if drawings is not None else 'ON_DEMAND','coordinate_space':'visible_rotated_page_normalized','confidence':round(conf,3),'quality_policy':'native-glyph-v2/ocr-confidence-v2','native_text_usable':good}
        content['ocr_pixel_cache_hit']=ocr_cache_hit
        return {'text':txt,'content':content,'quality':quality,'method':method,'seconds':round(time.perf_counter()-started,3)}

def preview_pdf(path):
    """Stable application pagination for DOCX/XML/XLSX. Never claim source Word pagination."""
    blocks=[]
    suffix=path.suffix.lower()
    if suffix=='.docx':
        d=Document(path)
        blocks.extend(p.text for p in d.paragraphs if p.text.strip())
        blocks.extend(' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows)
    elif suffix=='.xlsx':
        w=openpyxl.load_workbook(path,read_only=True,data_only=True)
        for s in w:
            blocks.append(s.title)
            blocks.extend(' | '.join(str(v) if v is not None else '' for v in r) for r in s.iter_rows(values_only=True))
        w.close()
    elif suffix=='.xml':
        root=ElementTree.parse(path).getroot()
        # EGRZ containers also carry CSS, images and multi-megabyte signatures.
        # Index the structured conclusion, not their binary/base64 payloads.
        conclusions=[e for e in root.iter() if e.tag.split('}')[-1]=='Conclusion']
        scopes=conclusions or [root]
        excluded={'stylesheet','style','script','Signature','signatures','qrcode',
                  'conclusionRaw','conclusionBody','X509Certificate','EncapsulatedCRLValue'}
        def xml_blocks(element):
            tag=element.tag.split('}')[-1]
            if tag in excluded:return
            text=norm(element.text or '')
            if text and not (len(text)>256 and re.fullmatch(r'[A-Za-z0-9+/=\s]+',text)):
                blocks.append(tag+': '+text)
            for child in element:xml_blocks(child)
        for scope in scopes:xml_blocks(scope)
    else:
        with Image.open(path) as img:
            out=fitz.open();page=out.new_page(width=img.width*.75,height=img.height*.75)
            page.insert_image(page.rect,filename=str(path));return out.tobytes()
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import textwrap
    font=ROOT/'assets'/'DejaVuSans.ttf'
    if not font.exists():font=Path('C:/Windows/Fonts/arial.ttf') if os.name=='nt' else Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    if 'Inspector' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('Inspector',str(font)))
    buffer=io.BytesIO();c=canvas.Canvas(buffer,pagesize=(842,595));y=560
    c.setFont('Inspector',10)
    for block in blocks:
        for line in textwrap.wrap(block,width=125,replace_whitespace=False) or ['']:
            if y<35:c.showPage();c.setFont('Inspector',10);y=560
            c.drawString(25,y,line);y-=14
        y-=5
    c.save();return buffer.getvalue()

def document_pdf(doc):
    source=ROOT/doc['path']
    if doc['format']=='.pdf':return source
    dest=DATA/'cache'/doc['sha256']/PARSER_VERSION/('preview-xml-v2.pdf' if doc['format']=='.xml' else 'preview.pdf')
    if not dest.exists():
        dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(preview_pdf(source))
    return dest

def parse_document(file_id,max_pages=0,allow_ocr=True,callback=None,cancel=None,page_numbers=None):
    doc=one('SELECT * FROM documents WHERE id=%s',(file_id,))
    if not doc:raise ValueError('Файл не найден')
    execute("UPDATE documents SET parse_status='PARSING' WHERE id=%s",(file_id,))
    pdf=None
    try:
        from .registry import sha256
        if sha256(inside_root(doc['path']))!=doc['sha256']:
            raise ValueError('SOURCE_CHANGED: содержимое исходника изменено после регистрации; загрузите новую редакцию')
        path=document_pdf(doc)
        pdf=fitz.open(path)
        if pdf.needs_pass:raise ValueError('PDF защищён паролем')
        total=len(pdf)
        execute('UPDATE documents SET page_count=%s WHERE id=%s',(total,file_id))
        processed_now=0
        for n in range(1,total+1):
            if page_numbers is not None and n not in page_numbers:continue
            if cancel and cancel():break
            old=one('SELECT id,quality,method,text FROM pages WHERE file_id=%s AND page=%s AND parser_version=%s',(file_id,n,PARSER_VERSION))
            needs_ocr=old and allow_ocr and old['method']=='text' and not native_text_usable(old['text'])
            if old and old['quality']!='ERROR' and not (old['quality']=='OCR_REQUIRED' and allow_ocr) and not needs_ocr:continue
            if max_pages and processed_now>=max_pages:break
            cache=DATA/'cache'/doc['sha256']/PARSER_VERSION/f'{n}-{int(allow_ocr)}.json'
            result=None
            if cache.exists():
                try:
                    cached=json.loads(cache.read_text(encoding='utf-8'))
                    if isinstance(cached,dict) and {'quality','text','content','method','seconds'}.issubset(cached) and isinstance(cached['content'].get('lines'),list):result=cached
                except (ValueError,OSError,AttributeError):pass # Interrupted cache is regenerated from the immutable source.
            if result is None or result['quality']=='ERROR' or (allow_ocr and result['method']=='text' and not native_text_usable(result['text'])):
                try:result=extract_page(path,n,allow_ocr,document=pdf)
                except Exception as e:result={'text':'','content':{'lines':[],'error':str(e)},'quality':'ERROR','method':'failed','seconds':0}
                write_cache(cache,result)
            # Some CAD PDFs expose NUL glyphs. PostgreSQL text/jsonb rejects them;
            # remove only that control code, including from pre-existing caches.
            def clean_nulls(value):
                if isinstance(value,str):return value.replace('\x00','')
                if isinstance(value,list):return [clean_nulls(v) for v in value]
                if isinstance(value,dict):return {k:clean_nulls(v) for k,v in value.items()}
                return value
            result=clean_nulls(result)
            if result['method'].startswith('tesseract') and result['content'].get('confidence',0)<.80:
                result['quality']='LOW_QUALITY'
            execute('''INSERT INTO pages(file_id,page,text,content,quality,method,seconds,parser_version) VALUES(%s,%s,%s,%s,%s,%s,%s,%s)
            ON CONFLICT(file_id,page) DO UPDATE SET text=excluded.text,content=excluded.content,quality=excluded.quality,method=excluded.method,seconds=excluded.seconds,parser_version=excluded.parser_version''',(file_id,n,result['text'],result['content'],result['quality'],result['method'],result['seconds'],PARSER_VERSION))
            execute('UPDATE documents SET parsed_pages=(SELECT count(*) FROM pages WHERE file_id=%s) WHERE id=%s',(file_id,file_id))
            processed_now+=1
            if callback:callback(n,total)
        counts=query('SELECT quality,count(*) AS count FROM pages WHERE file_id=%s GROUP BY quality',(file_id,))
        done=sum(r['count'] for r in counts)
        status='PARSED' if done==total and not any(r['quality'] in ('ERROR','OCR_REQUIRED') for r in counts) else 'PARTIAL'
        execute('UPDATE documents SET parse_status=%s,parsed_pages=%s,quality=%s WHERE id=%s',(status,done,{r['quality']:r['count'] for r in counts},file_id))
        return {'pages':total,'processed':done,'status':status,'quality':counts}
    except Exception as e:
        execute("UPDATE documents SET parse_status='ERROR',quality=%s WHERE id=%s",({'error':str(e)},file_id));raise
    finally:
        if pdf is not None:pdf.close()

NUMBER=r'[-+]?\d+(?:[ \u00a0]\d{3})*(?:[.,]\d+)?'
ENTITY=re.compile(r'(?:помещени[еяю]|пом\.|двер[ьи]|элемент|колонна|пилон|марка|ось|оси)\s*(?:№\s*)?([A-Za-zА-Яа-я]?\d[\w./-]{0,20})',re.I)

def object_label_valid(ordinal,quote):
    """Do not turn a similarly named site/room subtotal into a building indicator."""
    text=norm(quote).casefold()
    if ordinal==2 and re.search(r'общая\s+площадь\s+(?:работ|зоны|восстановления|травяного|газон|озеленения|территории|участка|помещения)',text):return False
    if ordinal==8 and re.search(r'высота\s+(?:здания\s+)?корпус\w*\s+\d',text):return False
    if ordinal==13:
        if re.search(r'актов\w*\s+зал|обеден\w*\s+зал|столов\w*|гардероб\w*|парков\w*|машино.?мест|бочк\w*|резервуар\w*|емкост\w*|ёмкост\w*',text):return False
        # A bare "capacity" may name a room or a container. Require object scope.
        if not re.search(r'школ\w*|учащ\w*|общеобразовательн\w*|детск\w*\s+сад|дошкольн\w*|вместимост\w*\s+(?:здания|объекта)|проектн\w*\s+мощност',text):return False
    return True

def scalar(raw,unit):
    raw=norm(raw)
    s=re.sub(r'\s+','',raw).replace(',','.').replace('−','-')
    if re.fullmatch(r'[-+]?\d+(?:\.\d+)?',s):
        v=float(s)
        if math.isfinite(v):return {'raw':raw,'normalized':v,'unit':unit,'kind':'number'}
    canonical=s.upper()
    if re.fullmatch(r'[АВCСAВB][\d.+-]*',canonical):canonical=canonical.translate(str.maketrans({'А':'A','В':'B','С':'C'}))
    return {'raw':raw,'normalized':canonical,'unit':unit,'kind':'text'}


def literal_unit(raw,quote):
    """Read the unit next to the source number; never accept an LLM unit over it."""
    pattern=r'(?<![\d.,])'+re.escape(raw)+r'\s*(Гкал/ч|м³/сут|м3/сут|м³/ч|м3/ч|мм|см|дм³|дм3|м²|м2|м³|м3|кВт|Вт|мин|м|%|‰)(?!\w)'
    units={m.group(1).replace('м2','м²').replace('м3','м³') for m in re.finditer(pattern,quote,re.I)}
    canonical={u.casefold():u for u in ('Гкал/ч','м³/сут','м³/ч','мм','см','дм³','м²','м³','кВт','Вт','мин','м','%','‰')}
    units={canonical.get(u.casefold(),u) for u in units}
    return next(iter(units)) if len(units)==1 else ('AMBIGUOUS' if units else None)

def datum_evidence(line,lines):
    """Building zero must be explicitly connected to an absolute elevation."""
    text=norm(line['text']);used=[line]
    if not re.search(r'0[.,]0{2,3}(?!\d)|нулев\w*\s+отмет',text,re.I):return None
    if 'абсолют' in text.lower() and not re.search(r'абсолют\w*\s+(?:высотн\w*\s+)?отмет\w*\s+[-+]?\d',text,re.I):
        b=line['bbox'];height=b[3]-b[1]
        following=[l for l in lines if 0<=l['bbox'][1]-b[3]<height*2.5 and abs(l['bbox'][0]-b[0])<.025 and re.match(r'отметк',norm(l['text']),re.I)]
        if following:
            other=min(following,key=lambda l:l['bbox'][1]);used.append(other);text+=' '+norm(other['text'])
    match=re.search(r'0[.,]0{2,3}\s*=\s*('+NUMBER+r')',text)
    if not match:match=re.search(r'абсолют\w*\s+(?:высотн\w*\s+)?отмет\w*\s*[:=]?\s*('+NUMBER+r')',text,re.I)
    if not match:return None
    return match.group(1),used

def extract_facts(file_id,parameters):
    doc=one('SELECT * FROM documents WHERE id=%s',(file_id,))
    pages=query('SELECT * FROM pages WHERE file_id=%s AND quality=%s ORDER BY page',(file_id,'GOOD'))
    touched=set()
    parameters=[p for p in parameters if p['config']['extractor']=='anchored']
    for page in pages:
        lines=page['content']['lines']
        for i,line in enumerate(lines):
            text=norm(line['text']).lower().replace('ё','е')
            for p in parameters:
                cfg=p['config']
                if cfg['extractor']!='anchored':continue
                datum=datum_evidence(line,lines) if p['ordinal']==9 else None
                if p['ordinal']==9 and datum is None:continue
                found=None
                for alias in cfg['aliases']:
                    if alias not in text:continue
                    m=re.search(re.escape(alias).replace(r'\ ',r'\s+'),text,re.I)
                    if m:found=m;break
                if not found and not datum:continue
                if datum:
                    # Morphological variants and a separate, aligned continuation line.
                    text=norm(' '.join(l['text'] for l in datum[1])).lower()
                    found=re.search(re.escape(datum[0]),text)
                if p['ordinal']==9 and re.search(r'верхн|кровл|максим',text[:found.start()]):continue
                tail=datum[0] if datum else text[found.end():].lstrip(' :—-=,;')
                tail=re.sub(r'^в\s*т\.\s*ч\.?\s*[:;,]?\s*','',tail)
                # Subtotals are not interchangeable with the building-wide parameter.
                if cfg['scope']=='OBJECT' and re.match(r'(?:наземн|надземн|подземн|жилой|жилых|техническ|автостоян|паркинг|без\s|с\s)',tail):continue
                used=datum[1] if datum else [line]
                # PDF tables often put the value on the same visual row in a separate line.
                if not re.search(r'\d',tail):
                    neighbors=[l for l in lines if l is not line and abs((l['bbox'][1]+l['bbox'][3])/2-(line['bbox'][1]+line['bbox'][3])/2)<.008 and l['bbox'][0]>=line['bbox'][2]-.01]
                    neighbors.sort(key=lambda l:l['bbox'][0])
                    if neighbors:
                        numeric=[l for l in neighbors if re.match(r'^'+NUMBER,l['text'])]
                        if any(re.search(NUMBER+r'\s*[-–/+×xх]\s*\d',l['text']) for l in numeric):continue
                        distinct={norm(l['text']) for l in numeric}
                        if len(distinct)>1:
                            label={'PD':'по пд','RD':'по рд','ID':'по ид'}.get(doc['stage'])
                            headers=[l for l in lines if label and label in norm(l['text']).lower() and l['bbox'][3]<line['bbox'][1]]
                            if headers:
                                header=max(headers,key=lambda l:l['bbox'][1]);hx=(header['bbox'][0]+header['bbox'][2])/2
                                numeric=sorted(numeric,key=lambda l:abs((l['bbox'][0]+l['bbox'][2])/2-hx))[:1]
                            else:continue
                        if not numeric:continue
                        chosen=numeric[0]
                        tail+=' '+chosen['text'];used.append(chosen)
                tail=tail.strip()
                typed_patterns={21:r'^[«"\s]*([A-GА-Г][+]{0,2})(?!\w)',124:r'^[«"\s]*([A-GА-Г][+]{0,2})(?!\w)',22:r'^([IVX]{1,4})(?!\w)',23:r'^([CС]\s*[0-3])(?!\d)',55:r'(?<!\w)([BВ]\s*\d+(?:[.,]\d+)?)(?!\w)',56:r'(?<!\w)([CС]\s*\d+)(?!\w)',57:r'(?<!\w)([AА]\s*\d+[CС]?)(?!\w)',103:r'(?<!\w)(EI\s*[- ]?\d+)(?!\w)'}
                match=re.search(typed_patterns[p['ordinal']],tail,re.I) if p['ordinal'] in typed_patterns else re.match(r'(?:[а-яa-z²³/%().·]+\s*[:—=-]?\s*){0,3}('+NUMBER+r')',tail,re.I)
                if not match:continue
                raw=match.group(1)
                if p['ordinal'] not in typed_patterns and re.match(r'\s*[-–/+×xх]\s*\d',tail[match.end():]):continue
                if p['ordinal']==7 and (re.match(r'\s*[,;]\s*\d',tail[match.end():]) or not float(raw.replace(' ','').replace(',','.')).is_integer()):continue
                em=ENTITY.search(line['text'])
                entity='OBJECT' if cfg['scope']=='OBJECT' else ('ELEMENT:'+em.group(1).upper() if em else None)
                if entity is None:continue # Do not silently compare unrelated elements.
                value=scalar(raw,p['unit'])
                if p['ordinal']==103:value.update(kind='number',normalized=float(re.search(r'\d+',raw).group()))
                # Explicit units take precedence; normalize supported metric conversions.
                explicit=re.search(re.escape(raw)+r'\s*(мм|см|м²|м2|м³|м3|м|кВт|Вт)(?!\w)',tail,re.I)
                if explicit and value['kind']=='number':
                    actual_unit=explicit.group(1).replace('м2','м²').replace('м3','м³')
                    factors={('мм','м'):.001,('см','м'):.01,('м','мм'):1000,('см','мм'):10,('Вт','кВт'):.001,('кВт','Вт'):1000}
                    if actual_unit!=p['unit']:
                        if (actual_unit,p['unit']) not in factors:continue
                        value['normalized']*=factors[(actual_unit,p['unit'])]
                    value['source_unit']=actual_unit
                box=[min(l['bbox'][0] for l in used),min(l['bbox'][1] for l in used),max(l['bbox'][2] for l in used),max(l['bbox'][3] for l in used)]
                quote=' '.join(l['text'] for l in used)
                if not object_label_valid(p['ordinal'],quote):continue
                evidence={'file_id':file_id,'sha256':doc['sha256'],'stage':doc['stage'],'document_code':doc['document_code'],'revision':doc['revision'],'approval_status':doc['approval_status'],'page':page['page'],'bbox':box,'quote':quote,'value':value,'entity':entity,'coordinate_space':'visible_rotated_page_normalized','source_pagination':'original' if doc['format']=='.pdf' else 'generated_preview'}
                method='anchored-datum-v5' if p['ordinal']==9 else 'anchored-v4'
                fact_id=uid(method,file_id,p['code'],entity,page['page'],box,raw)
                execute('INSERT INTO facts(id,file_id,object_id,parameter_code,entity,page,value,evidence,method,confidence) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(id) DO NOTHING',(fact_id,file_id,doc['object_id'],p['code'],entity,page['page'],value,evidence,method,line.get('confidence',1)))
                touched.add(p['code'])
    return sorted(touched)
